from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "see_e2f_rdb_wps_reproduction_final.pptx"
MANIFEST_PATH = OUT_DIR / "generation_manifest.json"
AUDIT_PATH = OUT_DIR / "text_layout_audit.json"
BRIEF_PATH = OUT_DIR / "brief.md"
CHECK_PATH = OUT_DIR / "check.md"

CANVAS_W = 1600
CANVAS_H = 900
SCALE = 100.0

BLUE = RGBColor(18, 61, 119)
BLUE2 = RGBColor(44, 89, 150)
LIGHT_BLUE = RGBColor(230, 240, 252)
GREEN = RGBColor(81, 132, 74)
LIGHT_GREEN = RGBColor(229, 240, 229)
ORANGE = RGBColor(222, 103, 24)
PURPLE = RGBColor(125, 91, 181)
RED = RGBColor(220, 26, 26)
GRAY = RGBColor(112, 112, 112)
LIGHT_GRAY = RGBColor(245, 247, 249)
INK = RGBColor(12, 31, 58)


audit_items: list[dict] = []


def emu(v: float):
    return Inches(v / SCALE)


def rgb_hex(c: RGBColor) -> str:
    return f"#{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def add_shape(slide, kind, x, y, w, h, fill=RGBColor(255, 255, 255), line=BLUE, radius=True, dash=None, lw=1.2):
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(lw)
    if dash:
        shape.line.dash_style = dash
    return shape


def set_text(shape, text, size=8, bold=False, color=INK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, italic=False):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    tf.margin_top = Pt(1.5)
    tf.margin_bottom = Pt(1.5)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = line
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    if hasattr(shape, "left"):
        audit_items.append(
            {
                "text": text,
                "x": round(shape.left / Inches(1) * SCALE, 2),
                "y": round(shape.top / Inches(1) * SCALE, 2),
                "w": round(shape.width / Inches(1) * SCALE, 2),
                "h": round(shape.height / Inches(1) * SCALE, 2),
                "font_size": size,
                "status": estimate_text_fit(text, shape.width / Inches(1) * 72, shape.height / Inches(1) * 72, size),
            }
        )


def estimate_text_fit(text: str, width_pt: float, height_pt: float, size: float) -> str:
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return "PASS"
    if all(len(ln) <= 2 for ln in lines):
        return "PASS"
    max_line = max(len(ln) for ln in lines)
    est_width = max_line * size * 0.52
    est_height = len(lines) * size * 1.18
    if est_width > width_pt * 0.96 or est_height > height_pt * 0.92:
        return "CAUTION"
    return "PASS"


def textbox(slide, x, y, w, h, text, size=8, bold=False, color=INK, align=PP_ALIGN.CENTER, italic=False):
    tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    set_text(tb, text, size=size, bold=bold, color=color, align=align, italic=italic)
    return tb


def panel(slide, x, y, w, h, title, color=BLUE):
    box = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=RGBColor(253, 255, 255), line=color, lw=1.4)
    textbox(slide, x + 8, y + 5, w - 16, 28, title, size=15, bold=True, color=INK)
    return box


def step_circle(slide, x, y, n, color=BLUE):
    c = add_shape(slide, MSO_SHAPE.OVAL, x, y, 27, 27, fill=color, line=color, lw=1)
    set_text(c, str(n), size=12, bold=True, color=RGBColor(255, 255, 255))
    return c


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.0, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = 3
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def line(slide, x1, y1, x2, y2, color=BLUE, width=1.0, dashed=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if dashed:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def small_button(slide, x, y, w, h, text, color=BLUE2, fill=LIGHT_BLUE, size=7.2):
    s = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=color, lw=0.8)
    set_text(s, text, size=size, color=INK)
    return s


def symbol(slide, x, y, kind, fill, line_color=None, size=13):
    line_color = line_color or fill
    if kind == "circle":
        shp = add_shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill=fill, line=line_color, lw=0.9)
    elif kind == "triangle":
        shp = add_shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, size, size, fill=fill, line=line_color, lw=0.9)
    elif kind == "diamond":
        shp = add_shape(slide, MSO_SHAPE.DIAMOND, x, y, size, size, fill=fill, line=line_color, lw=0.9)
    elif kind == "star":
        shp = add_shape(slide, MSO_SHAPE.STAR_5_POINT, x, y, size + 4, size + 4, fill=fill, line=line_color, lw=0.9)
    else:
        shp = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, size, size, fill=fill, line=line_color, lw=0.9)
    return shp


def mini_bars(slide, x, y, count=12, w=12, max_h=50, color=BLUE2, descending=True):
    for i in range(count):
        h = max_h * (0.9 - i * 0.055) if descending else max_h * (0.25 + i * 0.045)
        h = max(8, h)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + i * (w + 4), y + max_h - h, w, h, fill=color, line=color, lw=0.2)


def feature_strip(slide, x, y, count, colors):
    for i in range(count):
        c = colors[i % len(colors)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + i * 15, y, 13, 16, fill=c, line=RGBColor(210, 210, 210), lw=0.3)


def dashed_group(slide, x, y, w, h, line_color=RED):
    grp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=None, line=line_color, dash=MSO_LINE_DASH_STYLE.DASH, lw=1.2)
    return grp


def endpoint_cluster(slide, cx, cy, color, star_color):
    symbol(slide, cx - 10, cy - 34, "star", star_color, PURPLE, 15)
    pts = [(cx - 36, cy - 4), (cx - 16, cy + 16), (cx + 12, cy + 12), (cx + 34, cy - 3)]
    for px, py in pts:
        symbol(slide, px, py, "circle", RGBColor(238, 238, 248), color, 9)
        line(slide, cx, cy - 16, px + 4, py + 4, color=color, width=0.7)


def draw_left(slide):
    panel(slide, 15, 10, 435, 810, "SEE / E2F-Guided Initialization")
    inner_x, inner_w = 35, 395

    # Step 1
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 50, inner_w, 105, fill=RGBColor(252, 254, 255), line=BLUE2, lw=1.0)
    step_circle(slide, inner_x + 6, 58, 1)
    textbox(slide, inner_x + 60, 58, 210, 20, "HDLSS biomedical data", size=10, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, inner_x + 70, 82, 115, 28, "X ∈ Rⁿˣᵈ, y", size=10, italic=True)
    textbox(slide, inner_x + 30, 112, 120, 18, "Feature matrix  X", size=7)
    textbox(slide, inner_x + 245, 112, 105, 18, "Label vector  y", size=7)
    feature_strip(slide, inner_x + 28, 130, 12, [BLUE2, RGBColor(83, 148, 181), RGBColor(180, 210, 205), RGBColor(242, 223, 181)])
    textbox(slide, inner_x + 210, 126, 34, 18, "…", size=10)
    feature_strip(slide, inner_x + 255, 130, 8, [GREEN, RGBColor(178, 150, 101), RGBColor(232, 164, 83)])
    textbox(slide, inner_x + 372, 126, 20, 18, "…", size=10)
    arrow(slide, 235, 155, 235, 175, BLUE, 1.6)

    # Step 2
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 175, inner_w, 130, fill=RGBColor(252, 254, 255), line=BLUE2, lw=1.0)
    step_circle(slide, inner_x + 6, 183, 2)
    textbox(slide, inner_x + 54, 185, 230, 20, "Supervised relevance scoring", size=9.5, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, inner_x + 60, 210, 80, 18, "ANOVA F-score", size=7.4)
    textbox(slide, inner_x + 245, 210, 105, 18, "ranked relevance", size=7.2, italic=True)
    mini_bars(slide, inner_x + 28, 222, count=14, w=10, max_h=50, color=BLUE2)
    line(slide, inner_x + 25, 275, inner_x + 335, 275, color=INK, width=0.8)
    line(slide, inner_x + 25, 225, inner_x + 25, 275, color=INK, width=0.8)
    textbox(slide, inner_x + 320, 255, 36, 18, "…", size=10)
    small_button(slide, inner_x + 80, 285, 220, 22, "F-score  →  relevance order", size=7.5)
    arrow(slide, 235, 305, 235, 325, BLUE, 1.6)

    # Step 3
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 325, inner_w, 110, fill=RGBColor(252, 254, 255), line=BLUE2, lw=1.0)
    step_circle(slide, inner_x + 6, 333, 3)
    textbox(slide, inner_x + 60, 337, 250, 20, "Supervised candidate feature pool", size=10, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, inner_x + 60, 364, 130, 18, "Top-K supervised pool", size=7.2)
    textbox(slide, inner_x + 198, 364, 70, 18, "K ≤ 512", size=8.3, italic=True)
    for i, cx in enumerate([75, 110, 148, 186, 224, 262], start=1):
        symbol(slide, inner_x + cx - 10, 385, "circle", [RGBColor(230, 220, 247), LIGHT_BLUE, LIGHT_GREEN, RGBColor(221, 235, 211), RGBColor(230, 230, 230), RGBColor(224, 224, 224)][i - 1], INK, 20)
        textbox(slide, inner_x + cx - 9, 386, 18, 16, str(i), size=8, bold=True)
    textbox(slide, inner_x + 286, 386, 34, 16, "…", size=10)
    symbol(slide, inner_x + 326, 385, "circle", RGBColor(235, 235, 235), INK, 20)
    textbox(slide, inner_x + 328, 386, 16, 16, "K", size=8, bold=True)
    small_button(slide, inner_x + 30, 417, 135, 22, "poolBudget = 512", size=7.5)
    small_button(slide, inner_x + 190, 417, 135, 22, "high relevance", size=7.5)
    arrow(slide, 235, 435, 235, 455, BLUE, 1.6)

    # Red conceptual group
    dashed_group(slide, 18, 445, 425, 295)
    vlab = textbox(slide, -18, 505, 180, 18, "Key idea: Supervised E2F Shaping", size=6.7, bold=True, color=RED)
    vlab.rotation = 270

    # Step 4
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 445, inner_w, 155, fill=RGBColor(255, 253, 252), line=RGBColor(255, 255, 255), lw=0.2)
    step_circle(slide, inner_x + 6, 455, 4)
    textbox(slide, inner_x + 70, 455, 230, 20, "Redundancy-aware endpoint growth", size=9.2, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, inner_x + 74, 478, 250, 28, "Endpoint growth\nrelevance + redundancy + error", size=6.7, italic=True)
    endpoint_cluster(slide, inner_x + 58, 540, PURPLE, PURPLE)
    endpoint_cluster(slide, inner_x + 150, 540, GREEN, GREEN)
    endpoint_cluster(slide, inner_x + 240, 540, ORANGE, ORANGE)
    endpoint_cluster(slide, inner_x + 350, 540, BLUE2, BLUE2)
    textbox(slide, inner_x + 305, 518, 24, 15, "…", size=9)
    small_button(slide, inner_x + 14, 570, 84, 26, "|corr|\nredundancy", size=5.8)
    small_button(slide, inner_x + 120, 570, 86, 26, "single-feature\nKNN error", size=6.0)
    small_button(slide, inner_x + 230, 570, 80, 26, "branchWidth", size=6.5)
    small_button(slide, inner_x + 330, 570, 64, 26, "maxEntryK", size=6.5)
    arrow(slide, 235, 600, 235, 620, BLUE, 1.6)

    # Step 5
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 615, inner_w, 120, fill=RGBColor(255, 253, 252), line=RGBColor(255, 255, 255), lw=0.2)
    step_circle(slide, inner_x + 6, 625, 5)
    textbox(slide, inner_x + 125, 625, 190, 20, "Endpoint-to-Front Seeds", size=9.2, bold=True)
    textbox(slide, inner_x + 42, 655, 55, 30, "sparse\nendpoint", size=7.2)
    textbox(slide, inner_x + 328, 655, 60, 30, "low-error\nendpoint", size=7.2)
    symbol(slide, inner_x + 98, 655, "star", RGBColor(175, 145, 220), PURPLE, 17)
    pts = [(155, 675, PURPLE), (205, 692, GREEN), (245, 698, ORANGE), (308, 680, BLUE2)]
    for px, py, c in pts:
        symbol(slide, inner_x + px, py, "circle", RGBColor(238, 238, 248), c, 10)
    line(slide, inner_x + 110, 670, inner_x + 305, 700, PURPLE, 0.7, dashed=True)
    line(slide, inner_x + 140, 700, inner_x + 325, 678, GRAY, 0.7, dashed=True)
    textbox(slide, inner_x + 170, 655, 90, 18, "trade-off band", size=7.2, italic=True)
    small_button(slide, inner_x + 30, 725, 90, 28, "bandSteps", size=7)
    small_button(slide, inner_x + 160, 725, 100, 28, "bandMaxRatio", size=7)
    small_button(slide, inner_x + 300, 725, 95, 28, "multi-scale\nsubset size", size=6.2)
    arrow(slide, 235, 756, 235, 770, BLUE, 1.6)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, inner_x, 775, inner_w, 45, fill=RGBColor(245, 240, 255), line=PURPLE, lw=1.0)
    textbox(slide, inner_x + 20, 787, inner_w - 40, 20, "Initial CoverPop / SparsePop / EntryPop", size=9, bold=True)


def draw_middle(slide):
    panel(slide, 470, 10, 790, 810, "RDB-Guided Sparse Evolutionary Loop")
    # Entry arrow from left
    arrow(slide, 450, 405, 470, 405, BLUE, 2.2)

    # Population
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 485, 85, 180, 225, fill=RGBColor(247, 251, 255), line=BLUE2, lw=1.0)
    textbox(slide, 520, 100, 110, 24, "Population", size=10, bold=True)
    textbox(slide, 490, 130, 170, 20, "CoverPop + SparsePop + Archive", size=6.8)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 495, 165, 160, 90, fill=RGBColor(248, 249, 252), line=RGBColor(170, 184, 205), lw=0.8)
    for x, y, kind, c in [
        (510, 195, "circle", PURPLE), (525, 175, "circle", PURPLE), (540, 195, "circle", PURPLE),
        (555, 217, "triangle", BLUE2), (590, 176, "rect", GREEN), (606, 205, "rect", GREEN),
        (635, 178, "triangle", ORANGE), (630, 214, "triangle", ORANGE), (568, 210, "triangle", BLUE2),
    ]:
        symbol(slide, x, y, kind, RGBColor(235, 230, 250) if c == PURPLE else c, c, 12)
    textbox(slide, 630, 205, 40, 20, "…", size=9)
    small_button(slide, 500, 270, 48, 28, "NDSort", size=6.5)
    small_button(slide, 555, 270, 56, 28, "CrowdDis", size=6.5)
    small_button(slide, 615, 270, 40, 28, "FrontNo", size=6.0)

    # RDB core
    dashed_group(slide, 695, 65, 250, 240)
    textbox(slide, 720, 78, 200, 24, "RDB Core Controller", size=11, bold=True)
    textbox(slide, 715, 110, 210, 18, "Role-Decoupled Backbone (RDB)", size=8.2, bold=True)
    textbox(slide, 715, 136, 220, 20, "Role pressure ≠ fixed parent identity", size=6.8, italic=True)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 705, 170, 110, 68, fill=RGBColor(244, 240, 255), line=PURPLE, lw=0.9)
    set_text(slide.shapes[-1], "Extreme\nPreservation", size=8.5, bold=True)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 825, 170, 110, 68, fill=LIGHT_GREEN, line=GREEN, lw=0.9)
    set_text(slide.shapes[-1], "Coverage\nEnhancement", size=8.5, bold=True)
    small_button(slide, 710, 255, 100, 34, "extreme quota", color=PURPLE, fill=RGBColor(246, 242, 255), size=6.8)
    small_button(slide, 825, 255, 100, 34, "coverage quota", color=GREEN, fill=LIGHT_GREEN, size=6.8)

    # Offspring generator
    dashed_group(slide, 985, 65, 250, 370)
    textbox(slide, 995, 78, 232, 24, "Three-lane Offspring Generator", size=9.2, bold=True)
    textbox(slide, 1015, 105, 190, 18, "Cover / Sparse ≠ Band Offspring", size=7.8)
    # lanes
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 995, 125, 230, 84, fill=RGBColor(250, 245, 255), line=RGBColor(186, 160, 214), lw=0.8)
    textbox(slide, 1010, 132, 85, 18, "Cover lane", size=7.2, italic=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 156, 118, 18, "CoverOffDec", size=7.6, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 178, 125, 18, "general Pareto search", size=7.1, align=PP_ALIGN.LEFT)
    symbol(slide, 1158, 142, "star", RGBColor(185, 150, 225), PURPLE, 14)
    symbol(slide, 1186, 160, "star", RGBColor(185, 150, 225), PURPLE, 14)
    symbol(slide, 1158, 176, "star", RGBColor(185, 150, 225), PURPLE, 14)
    textbox(slide, 1205, 155, 25, 20, "…", size=9)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 995, 220, 230, 95, fill=RGBColor(251, 255, 251), line=RGBColor(180, 200, 180), lw=0.8)
    textbox(slide, 1010, 228, 85, 18, "Sparse lane", size=7.2, italic=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 250, 118, 18, "SparseOffDec", size=7.5, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 272, 105, 18, "Add / Drop / Swap", size=7.1, italic=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 294, 105, 18, "sparsity pressure", size=7.0, italic=True, align=PP_ALIGN.LEFT)
    for i in range(4):
        symbol(slide, 1135, 236 + i * 16, "circle", RGBColor(238, 238, 238), GRAY, 7)
        symbol(slide, 1152, 236 + i * 16, "circle", RGBColor(238, 238, 238), GRAY, 7)
    textbox(slide, 1164, 251, 28, 16, "↔", size=9)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1190, 235, 26, 62, fill=RGBColor(238, 248, 238), line=GREEN, lw=0.8)
    symbol(slide, 1198, 244, "circle", RGBColor(170, 110, 75), RGBColor(120, 80, 50), 6)
    symbol(slide, 1204, 270, "circle", RGBColor(238, 238, 238), GRAY, 6)
    line(slide, 1201, 247, 1206, 273, GRAY, 0.7)
    textbox(slide, 1205, 285, 25, 20, "…", size=9)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 995, 328, 230, 90, fill=RGBColor(248, 252, 255), line=RGBColor(170, 190, 210), lw=0.8)
    textbox(slide, 1010, 336, 85, 18, "Band lane", size=7.2, italic=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 358, 100, 18, "BandOffDec", size=7.6, bold=True, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 380, 130, 18, "GNNE²F Band Decoder", size=7.0, align=PP_ALIGN.LEFT)
    textbox(slide, 1010, 402, 115, 16, "front-band expansion", size=6.8, italic=True, align=PP_ALIGN.LEFT)
    # small curve
    prev = (1145, 342)
    for pt in [(1152, 358), (1165, 373), (1185, 385), (1212, 390)]:
        line(slide, prev[0], prev[1], pt[0], pt[1], PURPLE, 1.2)
        symbol(slide, pt[0] - 4, pt[1] - 4, "circle", RGBColor(235, 230, 248), PURPLE, 7)
        prev = pt
    line(slide, 1138, 338, 1148, 395, GRAY, 0.7)
    line(slide, 1140, 397, 1215, 405, GRAY, 0.7)
    textbox(slide, 1205, 340, 25, 20, "…", size=9)

    # Orange controls from RDB to Offspring
    arrow(slide, 945, 175, 985, 175, ORANGE, 2.2)
    arrow(slide, 945, 250, 985, 255, ORANGE, 2.2)

    # center evaluation stack
    arrow(slide, 820, 305, 820, 335, BLUE, 1.7)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 715, 345, 220, 105, fill=RGBColor(248, 253, 255), line=BLUE2, lw=1.0)
    textbox(slide, 745, 355, 160, 18, "Bi-objective Evaluation", size=7.8, bold=True)
    textbox(slide, 760, 377, 130, 17, "error + feature ratio", size=6.6, italic=True)
    small_button(slide, 725, 402, 95, 34, "classification\nerror", size=6.2)
    small_button(slide, 830, 402, 95, 34, "selected feature\nratio", color=GREEN, fill=LIGHT_GREEN, size=6.2)
    arrow(slide, 820, 450, 820, 475, BLUE, 1.7)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 715, 475, 220, 55, fill=RGBColor(248, 253, 255), line=BLUE2, lw=1.0)
    textbox(slide, 742, 484, 165, 20, "Environmental Selection", size=7.8, bold=True)
    textbox(slide, 777, 507, 100, 18, "NDSort + CrowdDis", size=7.2)
    arrow(slide, 820, 530, 820, 560, BLUE, 1.7)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 715, 560, 220, 100, fill=RGBColor(248, 253, 255), line=BLUE2, lw=1.0)
    textbox(slide, 750, 572, 150, 18, "Sparse Pareto Archive", size=8.0, bold=True)
    textbox(slide, 785, 594, 80, 18, "archive update", size=6.8, italic=True)
    for x, y, kind, c in [(735, 620, "circle", PURPLE), (755, 637, "circle", PURPLE), (790, 620, "rect", GREEN), (808, 640, "rect", GREEN), (855, 630, "triangle", ORANGE), (885, 620, "triangle", ORANGE), (922, 632, "triangle", BLUE2)]:
        symbol(slide, x, y, kind, RGBColor(237, 232, 249) if c == PURPLE else c, c, 12)
    textbox(slide, 907, 615, 30, 18, "…", size=9)

    # feedback loops
    line(slide, 715, 610, 565, 610, BLUE, 2.4)
    line(slide, 565, 610, 565, 310, BLUE, 2.4)
    arrow(slide, 565, 310, 565, 310, BLUE, 2.4)
    line(slide, 935, 610, 1120, 610, BLUE, 2.4)
    line(slide, 1120, 610, 1120, 435, BLUE, 2.4)
    arrow(slide, 1120, 435, 1120, 435, BLUE, 2.4)
    line(slide, 635, 315, 715, 410, GRAY, 1.0, dashed=True)
    line(slide, 945, 410, 995, 380, GRAY, 1.0, dashed=True)
    line(slide, 685, 300, 685, 480, GRAY, 1.0, dashed=True)
    arrow(slide, 685, 480, 715, 500, GRAY, 1.0, dashed=True)

    # Front formation
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 525, 675, 595, 145, fill=RGBColor(255, 255, 255), line=BLUE2, lw=1.0)
    textbox(slide, 770, 683, 130, 20, "Front Formation", size=9.2, bold=True)
    textbox(slide, 550, 690, 80, 35, "sparse\nendpoint", size=7.2)
    textbox(slide, 1030, 690, 80, 35, "low-error\nendpoint", size=7.2)
    symbol(slide, 612, 715, "star", RGBColor(185, 150, 225), PURPLE, 18)
    symbol(slide, 1058, 735, "triangle", BLUE2, BLUE2, 16)
    pts = [(650, 738, PURPLE), (720, 760, GREEN), (795, 773, GREEN), (900, 758, ORANGE), (960, 754, ORANGE), (1030, 740, BLUE2)]
    prev = (620, 728)
    for px, py, c in pts:
        line(slide, prev[0], prev[1], px, py, GRAY, 0.7)
        symbol(slide, px - 5, py - 5, "circle", RGBColor(235, 240, 248), c, 10)
        prev = (px, py)
    line(slide, 610, 750, 715, 790, PURPLE, 0.8, dashed=True)
    line(slide, 840, 780, 1060, 750, BLUE2, 0.8, dashed=True)
    arrow(slide, 680, 732, 690, 732, PURPLE, 1.0)
    arrow(slide, 1040, 710, 960, 720, BLUE2, 1.0)
    textbox(slide, 785, 725, 95, 18, "trade-off band", size=7.2, italic=True)
    textbox(slide, 610, 794, 90, 17, "continuity", size=6.5)
    textbox(slide, 790, 800, 80, 17, "sparsity", size=6.5)
    textbox(slide, 990, 794, 80, 17, "coverage", size=6.5)

    # Exit arrow to right panel
    arrow(slide, 1260, 405, 1280, 405, BLUE, 2.2)


def draw_right(slide):
    panel(slide, 1280, 10, 305, 810, "Final Output")
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1330, 75, 205, 72, fill=LIGHT_BLUE, line=BLUE2, lw=1.0)
    textbox(slide, 1360, 90, 145, 20, "Termination Criteria", size=9, bold=True)
    textbox(slide, 1375, 118, 120, 18, "maxFE / convergence", size=7.5)
    arrow(slide, 1430, 150, 1430, 225, BLUE, 2.0)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1320, 230, 230, 90, fill=LIGHT_BLUE, line=BLUE2, lw=1.0)
    textbox(slide, 1330, 250, 210, 50, "Stable Sparse\nPareto Feature Panels", size=12, bold=True)

    textbox(slide, 1360, 385, 125, 28, "Pareto Front", size=11, bold=True)
    # chart box area
    x0, y0 = 1322, 620
    # axes
    line(slide, 1330, 620, 1560, 620, INK, 1.2)
    line(slide, 1330, 620, 1330, 405, INK, 1.2)
    arrow(slide, 1330, 620, 1560, 620, INK, 1.2)
    arrow(slide, 1330, 620, 1330, 405, INK, 1.2)
    # grid
    for gx in [1390, 1440, 1490, 1540]:
        line(slide, gx, 615, gx, 420, RGBColor(200, 205, 210), 0.6, dashed=True)
    for gy in [560, 515, 470, 425]:
        line(slide, 1340, gy, 1560, gy, RGBColor(200, 205, 210), 0.6, dashed=True)
    # shaded band / curves approximated
    line(slide, 1350, 425, 1380, 495, RGBColor(210, 225, 242), 10)
    line(slide, 1380, 495, 1410, 545, RGBColor(210, 225, 242), 10)
    line(slide, 1410, 545, 1460, 580, RGBColor(210, 225, 242), 10)
    line(slide, 1460, 580, 1540, 605, RGBColor(210, 225, 242), 10)
    pts = [(1348, 425), (1365, 465), (1390, 515), (1430, 560), (1475, 585), (1515, 602)]
    for a, b in zip(pts, pts[1:]):
        line(slide, a[0], a[1], b[0], b[1], BLUE, 2.0)
    for px, py in pts:
        symbol(slide, px - 6, py - 6, "circle", RGBColor(245, 248, 255), BLUE, 12)
    axis_y = textbox(slide, 1278, 490, 130, 20, "Classification Error", size=7.6, bold=True)
    axis_y.rotation = 270
    textbox(slide, 1378, 635, 150, 20, "Selected Feature Ratio", size=8.2, bold=True)

    for x, text in [(1300, "HV"), (1390, "IGD"), (1474, "Stability")]:
        small_button(slide, x, 686, 78, 38, text, size=8.5)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1298, 748, 265, 50, fill=LIGHT_BLUE, line=BLUE2, lw=1.0)
    textbox(slide, 1342, 760, 178, 24, "biomedical feature panel", size=8.5, italic=True)


def draw_legends(slide):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 15, 842, 965, 48, fill=RGBColor(255, 255, 255), line=RGBColor(185, 190, 198), lw=0.8)
    legend = [
        (45, "star", PURPLE, "Extreme Anchor"),
        (205, "diamond", GREEN, "Diverse Anchor"),
        (360, "rect", RGBColor(232, 164, 83), "Medium Anchor"),
        (505, "triangle", BLUE2, "Compact Anchor"),
        (680, "circle", RGBColor(255, 255, 255), "Candidate Solution"),
        (845, "circle", RGBColor(255, 255, 255), "Selected (Survivor)"),
    ]
    for x, kind, c, label in legend:
        symbol(slide, x, 858, kind, c, INK if c == RGBColor(255, 255, 255) else c, 15)
        textbox(slide, x + 30, 854, 120, 22, label, size=7.2, align=PP_ALIGN.LEFT)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1000, 842, 585, 48, fill=RGBColor(255, 255, 255), line=RGBColor(185, 190, 198), lw=0.8)
    arrow(slide, 1020, 866, 1065, 866, BLUE, 1.8)
    textbox(slide, 1080, 854, 130, 22, "Data / Evolution Flow", size=7.1, align=PP_ALIGN.LEFT)
    arrow(slide, 1210, 866, 1258, 866, ORANGE, 1.8)
    textbox(slide, 1272, 854, 135, 22, "Control / Role Guidance", size=7.1, align=PP_ALIGN.LEFT)
    line(slide, 1415, 866, 1465, 866, GRAY, 1.6, dashed=True)
    arrow(slide, 1455, 866, 1465, 866, GRAY, 1.6, dashed=True)
    textbox(slide, 1475, 854, 95, 22, "Auxiliary Link", size=7.1, align=PP_ALIGN.LEFT)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(252, 253, 252)

    # subtle page frame
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, CANVAS_W, CANVAS_H, fill=None, line=RGBColor(245, 245, 245), lw=0.2)
    draw_left(slide)
    draw_middle(slide)
    draw_right(slide)
    draw_legends(slide)

    prs.save(PPTX_PATH)

    cautions = [it for it in audit_items if it["status"] != "PASS"]
    AUDIT_PATH.write_text(
        json.dumps(
            {
                "canvas": {"width": CANVAS_W, "height": CANVAS_H, "unit": "ppt-coordinate"},
                "total_text_items": len(audit_items),
                "caution_count": len(cautions),
                "cautions": cautions,
                "items": audit_items,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    MANIFEST_PATH.write_text(
        json.dumps(
            {
                "figure": "SEE / E2F + RDB-guided sparse evolutionary loop reproduction",
                "delivery_level": "L3-ish WPS/PPT native editable shapes; plots are editable line/shape approximations",
                "source": str(PPTX_PATH),
                "audit": str(AUDIT_PATH),
                "style": {
                    "font": "Arial",
                    "primary_blue": rgb_hex(BLUE),
                    "control_orange": rgb_hex(ORANGE),
                    "dashed_red": rgb_hex(RED),
                },
                "strict_boundaries": [
                    "No numeric scientific data inferred beyond labels visible in the reference image.",
                    "Pareto and mini-plots are schematic shape reproductions, not measured curves.",
                    "All components are generated as editable PPT shapes where supported by PowerPoint/WPS.",
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    BRIEF_PATH.write_text(
        """# SEE / E2F + RDB WPS 复现 Brief

图型：科研算法流程 / 架构机制复现图
用途：论文/汇报插图，可在 WPS 演示中打开并编辑主要形状与文字
视觉目标：结构、主面板、颜色语义、箭头方向、关键标签、图例和整体比例接近参考图
编辑目标：L3-ish WPS/PPT 原生形状版；复杂曲线和散点用可编辑形状近似，不声明像素级逐元素完全一致
事实边界：只复现参考图中可见标签，不补充未经给出的指标数值或实验结论
验收标准：按 quality-gates Gate 4.1/4.2/4.3/4.4 严格检查，无文字溢出、重叠、错箭头、目标软件导出错版
""",
        encoding="utf-8",
    )
    CHECK_PATH.write_text(
        """# SEE / E2F + RDB WPS 复现验收记录

## 复现目标
- 视觉目标：结构和主视觉接近参考图，不承诺像素级完全一致。
- 编辑目标：WPS/PPT 原生形状可编辑，适合继续改文字、移动模块和调整箭头。
- 交付等级：L3-ish；小型曲线和示意图用 Office 线条/形状近似。

## 资产来源
- 图标：PPT/WPS 原生圆、三角、菱形、星形、矩形、箭头、线条；无字母占位冒充图标。
- 字体：Arial。
- 数据/标签：来自用户提供参考图可见文本；曲线和散点为示意复现，不代表真实数据。

## 自动文本审计
- 见 `text_layout_audit.json`。

## 目标软件验收
- WPS：待 WPS MCP 打开并导出后补充。
- 导出 PNG：待 WPS MCP 转换后补充。

## 全图裁剪证据
- 待基于 WPS 导出图片生成 contact sheet。

## 残余风险
- 参考图中极小 OCR 文本存在少量不可辨识处，已按可见语义复现，不额外发明。
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    build()
    print(PPTX_PATH)
