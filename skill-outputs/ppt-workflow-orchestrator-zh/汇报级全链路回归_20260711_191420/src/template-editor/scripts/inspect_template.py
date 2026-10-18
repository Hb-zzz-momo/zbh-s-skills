"""Inspect presentation structure using python-pptx."""
from __future__ import annotations
import argparse, json
from collections import Counter
from pathlib import Path
from pptx import Presentation

def shape_text(shape) -> str:
    return shape.text.strip() if getattr(shape, "has_text_frame", False) else ""

def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect PPTX slide size, layouts, text, shapes, images, tables, charts and notes.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path, required=True)
    parser.add_argument("--markdown", type=Path)
    args = parser.parse_args()
    source = args.pptx.resolve()
    if not source.is_file(): raise SystemExit(f"PPTX not found: {source}")
    prs = Presentation(source)
    slides=[]; layout_counts=Counter()
    for index, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout.name or "(unnamed)"; layout_counts[layout]+=1
        texts=[t for shape in slide.shapes if (t:=shape_text(shape))]
        fonts=sorted({run.font.name for shape in slide.shapes if getattr(shape,"has_text_frame",False) for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.name})
        slides.append({"slide":index,"layout":layout,"shape_count":len(slide.shapes),"pictures":sum(1 for s in slide.shapes if s.shape_type==13),"tables":sum(1 for s in slide.shapes if getattr(s,"has_table",False)),"charts":sum(1 for s in slide.shapes if getattr(s,"has_chart",False)),"texts":texts,"fonts":fonts,"notes":slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""})
    result={"source":str(source),"slide_size_inches":{"width":prs.slide_width/914400,"height":prs.slide_height/914400},"slide_count":len(slides),"layouts":dict(layout_counts),"slides":slides}
    json_path=args.json_path.resolve(); json_path.parent.mkdir(parents=True,exist_ok=True); json_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    if args.markdown:
        lines=["# 模板分析", "", f"- 文件：`{source}`", f"- 页数：{len(slides)}", f"- 尺寸：{result['slide_size_inches']['width']:.3f} × {result['slide_size_inches']['height']:.3f} in", "", "## 页面", "", "| 页 | Layout | Shapes | 图片 | 表格 | 图表 |", "|---:|---|---:|---:|---:|---:|"]
        lines += [f"| {s['slide']} | {s['layout']} | {s['shape_count']} | {s['pictures']} | {s['tables']} | {s['charts']} |" for s in slides]
        md=args.markdown.resolve(); md.parent.mkdir(parents=True,exist_ok=True); md.write_text("\n".join(lines)+"\n",encoding="utf-8")
    print(json.dumps({"slides":len(slides),"json":str(json_path)},ensure_ascii=False))
    return 0

if __name__ == "__main__": raise SystemExit(main())

