"""Check native PPTX shape bounds using python-pptx."""
from __future__ import annotations
import argparse, json
from pathlib import Path
from pptx import Presentation

EMU=914400
def main() -> int:
    parser=argparse.ArgumentParser(description="Check that every slide shape stays within the slide canvas.")
    parser.add_argument("pptx",type=Path); parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args(); source=args.pptx.resolve()
    if not source.is_file(): raise SystemExit(f"PPTX not found: {source}")
    prs=Presentation(source); width=prs.slide_width/EMU; height=prs.slide_height/EMU; issues=[]; checked=0
    for si,slide in enumerate(prs.slides,1):
        text_boxes=[]
        for shape in slide.shapes:
            checked+=1; x,y,w,h=[v/EMU for v in (shape.left,shape.top,shape.width,shape.height)]
            if x < -0.001 or y < -0.001 or x+w > width+0.001 or y+h > height+0.001: issues.append({"severity":"ERROR","slide":si,"shape_id":shape.shape_id,"name":shape.name,"message":"shape_out_of_bounds","box":{"x":x,"y":y,"w":w,"h":h}})
            # Slide-number placeholders intentionally share the footer band and
            # are not content boxes. Including them creates a false overlap on
            # every slide with a source footer.
            if getattr(shape,"has_text_frame",False) and shape.text.strip() and "Slide Number Placeholder" not in shape.name: text_boxes.append((shape.shape_id,shape.name,x,y,w,h))
        for index,a in enumerate(text_boxes):
            for b in text_boxes[index+1:]:
                ix=max(0,min(a[2]+a[4],b[2]+b[4])-max(a[2],b[2])); iy=max(0,min(a[3]+a[5],b[3]+b[5])-max(a[3],b[3])); overlap=ix*iy; smaller=min(a[4]*a[5],b[4]*b[5])
                if smaller and overlap/smaller > 0.20: issues.append({"severity":"ERROR","slide":si,"shape_ids":[a[0],b[0]],"names":[a[1],b[1]],"message":"text_boxes_overlap","overlap_ratio":round(overlap/smaller,3)})
    result={"source":str(source),"slide_size":{"width":width,"height":height},"checked_shapes":checked,"issues":issues,"status":"FAIL" if issues else "PASS"}; output=args.output.resolve(); output.parent.mkdir(parents=True,exist_ok=True); output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps(result,ensure_ascii=False)); return 2 if issues else 0

if __name__=="__main__": raise SystemExit(main())
