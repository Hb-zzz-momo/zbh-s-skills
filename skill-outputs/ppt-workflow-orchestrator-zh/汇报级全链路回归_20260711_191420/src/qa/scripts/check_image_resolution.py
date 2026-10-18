"""Estimate effective PPI of raster images placed in a PPTX."""
from __future__ import annotations
import argparse, io, json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU=914400
def main() -> int:
    parser=argparse.ArgumentParser(description="Check effective image PPI; default warning threshold is 150 PPI.")
    parser.add_argument("pptx",type=Path); parser.add_argument("--output",type=Path,required=True); parser.add_argument("--min-ppi",type=float,default=150)
    args=parser.parse_args(); source=args.pptx.resolve()
    if not source.is_file(): raise SystemExit(f"PPTX not found: {source}")
    prs=Presentation(source); images=[]; issues=[]
    for si,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE: continue
            # PptxGenJS stores an SVG plus a small PNG compatibility preview.
            # python-pptx exposes the preview as the picture blob; PPI is not a
            # meaningful quality metric for the vector object.
            if shape.name.lower().endswith("-svg"):
                images.append({"slide":si,"shape_id":shape.shape_id,"name":shape.name,"kind":"vector","effective_ppi":None})
                continue
            with Image.open(io.BytesIO(shape.image.blob)) as image: px_w,px_h=image.size
            in_w=shape.width/EMU; in_h=shape.height/EMU; ppi=min(px_w/in_w,px_h/in_h) if in_w and in_h else 0
            item={"slide":si,"shape_id":shape.shape_id,"pixels":[px_w,px_h],"size_inches":[in_w,in_h],"effective_ppi":round(ppi,1)}; images.append(item)
            if ppi < args.min_ppi: issues.append({"severity":"WARNING",**item,"message":"low_effective_ppi"})
    result={"source":str(source),"threshold_ppi":args.min_ppi,"images":images,"issues":issues,"status":"CAUTION" if issues else "PASS"}; output=args.output.resolve(); output.parent.mkdir(parents=True,exist_ok=True); output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps(result,ensure_ascii=False)); return 0

if __name__=="__main__": raise SystemExit(main())
