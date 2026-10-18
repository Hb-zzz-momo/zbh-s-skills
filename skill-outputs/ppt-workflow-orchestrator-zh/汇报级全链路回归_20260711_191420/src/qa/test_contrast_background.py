import subprocess, sys, tempfile, unittest
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
SCRIPT=ROOT/"scripts"/"check_contrast.py"

class ContrastBackgroundTests(unittest.TestCase):
    def test_resolves_slide_and_underlay_backgrounds(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); deck=td/"deck.pptx"; out=td/"contrast.json"; prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(0x19,0x3C,0x34)
            box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)); run=box.text_frame.paragraphs[0].add_run(); run.text="WHITE ON DARK"; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(20); run.font.bold=True
            panel=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(3),Inches(4),Inches(1)); panel.fill.solid(); panel.fill.fore_color.rgb=RGBColor(0x19,0x3C,0x34); panel.line.fill.background()
            label=slide.shapes.add_textbox(Inches(1.2),Inches(3.2),Inches(3.5),Inches(.5)); run=label.text_frame.paragraphs[0].add_run(); run.text="UNDERLAY"; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(14)
            prs.save(deck); cp=subprocess.run([sys.executable,str(SCRIPT),str(deck),"--output",str(out)],text=True,capture_output=True); self.assertEqual(cp.returncode,0,cp.stdout+cp.stderr)

if __name__=="__main__": unittest.main()
