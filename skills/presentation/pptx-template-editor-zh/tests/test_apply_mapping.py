import json, subprocess, sys, tempfile, unittest, zipfile
from pathlib import Path
from pptx import Presentation
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
SCRIPT=ROOT/"scripts"/"apply_mapping.py"

class ApplyMappingTests(unittest.TestCase):
    def test_replaces_expected_text_without_touching_source(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); source=td/"source.pptx"; output=td/"edited.pptx"
            old_img=td/"old.png"; new_img=td/"new.png"; Image.new("RGB",(40,40),"red").save(old_img); Image.new("RGB",(40,40),"green").save(new_img)
            prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_textbox(100000,100000,3000000,500000).text="OLD"; picture=slide.shapes.add_picture(str(old_img),100000,800000,800000,800000); picture.name="Hero Image"; prs.save(source)
            before=source.read_bytes()
            mapping={"operations":[{"type":"replace_text","slide":1,"expected":"OLD","replacement":"NEW"},{"type":"replace_image","slide":1,"shape_name":"Hero Image","replacement":str(new_img)}]}
            (td/"mapping.json").write_text(json.dumps(mapping),encoding="utf-8")
            cp=subprocess.run([sys.executable,str(SCRIPT),str(source),"--mapping",str(td/"mapping.json"),"--output",str(output),"--report",str(td/"report.json")],text=True,capture_output=True)
            self.assertEqual(cp.returncode,0,cp.stdout+cp.stderr)
            self.assertEqual(source.read_bytes(),before)
            edited=Presentation(output); self.assertIn("NEW",[s.text for s in edited.slides[0].shapes if getattr(s,"has_text_frame",False)])
            self.assertEqual(next(s for s in edited.slides[0].shapes if s.name=="Hero Image").image.blob,new_img.read_bytes())

    def test_rejects_zip_path_traversal(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); source=td/"bad.pptx"
            with zipfile.ZipFile(source,"w") as z: z.writestr("../escape.txt","bad")
            (td/"mapping.json").write_text('{"operations":[]}',encoding="utf-8")
            cp=subprocess.run([sys.executable,str(SCRIPT),str(source),"--mapping",str(td/"mapping.json"),"--output",str(td/"out.pptx"),"--report",str(td/"report.json")],text=True,capture_output=True)
            self.assertNotEqual(cp.returncode,0); self.assertIn("unsafe ZIP member",cp.stdout+cp.stderr)

if __name__ == "__main__": unittest.main()
