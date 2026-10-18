"""Check native chart existence and common readability risks."""
import argparse, json, zipfile
from pathlib import Path
from pptx import Presentation
def main():
    p=argparse.ArgumentParser(); p.add_argument("pptx",type=Path); p.add_argument("--output",type=Path,required=True); p.add_argument("--min-charts",type=int,default=0); a=p.parse_args(); prs=Presentation(a.pptx); charts=[]; issues=[]
    for si,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if getattr(shape,"has_chart",False): charts.append({"slide":si,"shape_id":shape.shape_id,"series":len(shape.chart.series),"categories":len(list(shape.chart.plots[0].categories)) if shape.chart.plots else 0})
    with zipfile.ZipFile(a.pptx) as z: xml="\n".join(z.read(n).decode("utf-8",errors="ignore") for n in z.namelist() if n.startswith("ppt/charts/chart") and n.endswith(".xml"))
    if len(charts)<a.min_charts: issues.append({"severity":"ERROR","message":f"native_chart_count {len(charts)} < {a.min_charts}"})
    if "3DChart" in xml: issues.append({"severity":"ERROR","message":"3D chart is forbidden"})
    for c in charts:
        if c["series"]==0 or c["categories"]==0: issues.append({"severity":"ERROR",**c,"message":"empty chart"})
        if c["categories"]>12: issues.append({"severity":"WARNING",**c,"message":"too many categories for slide readability"})
    result={"status":"BLOCK" if any(i["severity"]=="ERROR" for i in issues) else "CAUTION" if issues else "PASS","charts":charts,"issues":issues}; a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); print(json.dumps(result,ensure_ascii=False)); return 2 if result["status"]=="BLOCK" else 0
if __name__=="__main__": raise SystemExit(main())
