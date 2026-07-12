"""Check direct text/fill color contrast where both colors are explicitly resolvable."""
import argparse, json
from pathlib import Path
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.oxml.ns import qn

def rgb(hexv): return tuple(int(hexv[i:i+2],16)/255 for i in (0,2,4))
def lum(hexv):
    return sum(c*(v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4) for c,v in zip((.2126,.7152,.0722),rgb(hexv)))
def ratio(a,b):
    x,y=sorted((lum(a),lum(b)),reverse=True); return (x+.05)/(y+.05)
def blend(foreground,background,opacity):
    fg=tuple(int(foreground[i:i+2],16) for i in (0,2,4)); bg=tuple(int(background[i:i+2],16) for i in (0,2,4)); return "".join(f"{round(f*opacity+b*(1-opacity)):02X}" for f,b in zip(fg,bg))
def direct_color(color):
    try: return str(color.rgb) if color and color.type and color.rgb else None
    except Exception: return None
def contains(outer, inner, tolerance=1000):
    return outer.left<=inner.left+tolerance and outer.top<=inner.top+tolerance and outer.left+outer.width>=inner.left+inner.width-tolerance and outer.top+outer.height>=inner.top+inner.height-tolerance
def fill_style(shape):
    fill=getattr(shape,"fill",None)
    if fill is None or fill.type!=MSO_FILL_TYPE.SOLID: return None
    color=direct_color(fill.fore_color)
    if not color: return None
    opacity=1.0
    try:
        solid=shape._element.spPr.find(qn("a:solidFill")); color_node=next(iter(solid)) if solid is not None and len(solid) else None; alpha=color_node.find(qn("a:alpha")) if color_node is not None else None
        if alpha is not None: opacity=int(alpha.get("val"))/100000
    except Exception: pass
    return color,opacity
def background_for(slide, shapes, index, shape, fallback):
    slide_bg=fallback
    try:
        fill=slide.background.fill
        if fill.type==MSO_FILL_TYPE.SOLID: slide_bg=direct_color(fill.fore_color) or fallback
    except Exception: pass
    lower=slide_bg
    # Text boxes are commonly placed over a separate filled rectangle. Resolve
    # the closest lower z-order solid shape that fully contains the text box.
    for under in reversed(shapes[:index]):
        style=fill_style(under)
        if style and contains(under,shape):
            lower=blend(style[0],slide_bg,style[1]); break
    own=fill_style(shape)
    return blend(own[0],lower,own[1]) if own else lower
def main():
    p=argparse.ArgumentParser(); p.add_argument("pptx",type=Path); p.add_argument("--output",type=Path,required=True); p.add_argument("--min-ratio",type=float,default=4.5); p.add_argument("--background",default="FFFFFF"); a=p.parse_args(); prs=Presentation(a.pptx); issues=[]; checked=0; unresolved=0
    for si,slide in enumerate(prs.slides,1):
        shapes=list(slide.shapes)
        for index,shape in enumerate(shapes):
            if not getattr(shape,"has_text_frame",False) or not shape.text.strip(): continue
            bg=background_for(slide,shapes,index,shape,a.background)
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    fg=direct_color(run.font.color)
                    if not fg or not bg: unresolved+=1; continue
                    checked+=1; r=ratio(fg,bg); size=run.font.size.pt if run.font.size else 0; threshold=3.0 if size>=24 or (size>=18 and run.font.bold) else a.min_ratio
                    if r<threshold: issues.append({"severity":"ERROR","slide":si,"shape_id":shape.shape_id,"message":"low_text_contrast","ratio":round(r,2),"threshold":threshold,"font_size":size,"bold":bool(run.font.bold),"foreground":fg,"background":bg})
    result={"status":"BLOCK" if issues else "PASS","checked_runs":checked,"unresolved_runs":unresolved,"issues":issues}; a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); print(json.dumps(result,ensure_ascii=False)); return 2 if issues else 0
if __name__=="__main__": raise SystemExit(main())
