"""Extract PPTX text and notes with markitdown when available, otherwise python-pptx."""
from __future__ import annotations
import argparse, json, re
from pathlib import Path
from pptx import Presentation

PATTERN=re.compile(r"lorem|ipsum|xxxx|待替换|placeholder",re.I)

def main() -> int:
    parser=argparse.ArgumentParser(description="Extract slide text and notes; report placeholder-like content.")
    parser.add_argument("pptx",type=Path); parser.add_argument("--markdown",type=Path,required=True); parser.add_argument("--json",dest="json_path",type=Path,required=True)
    args=parser.parse_args(); source=args.pptx.resolve()
    if not source.is_file(): raise SystemExit(f"PPTX not found: {source}")
    prs=Presentation(source); slides=[]; placeholders=[]
    for index,slide in enumerate(prs.slides,1):
        texts=[shape.text.strip() for shape in slide.shapes if getattr(shape,"has_text_frame",False) and shape.text.strip()]; notes=slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        for value in [*texts,notes]:
            if value and PATTERN.search(value): placeholders.append({"slide":index,"text":value})
        slides.append({"slide":index,"texts":texts,"notes":notes})
    engine="python-pptx"; markdown_text=""
    try:
        from markitdown import MarkItDown
        markdown_text=MarkItDown().convert(source).text_content
        engine="markitdown+pptx"
    except Exception:
        markdown_text=""
    result={"source":str(source),"engine":engine,"slide_count":len(slides),"slides":slides,"placeholder_hits":placeholders,"status":"FAIL" if placeholders else "PASS"}
    json_path=args.json_path.resolve(); json_path.parent.mkdir(parents=True,exist_ok=True); json_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    if not markdown_text:
        lines=["# PPT 文本提取",""]
        for slide in slides: lines += [f"## Slide {slide['slide']}","",*slide["texts"],"",f"备注：{slide['notes']}" if slide["notes"] else "","" ]
        markdown_text="\n".join(lines)
    md=args.markdown.resolve(); md.parent.mkdir(parents=True,exist_ok=True); md.write_text(markdown_text,encoding="utf-8"); print(json.dumps({"engine":engine,"slides":len(slides),"placeholders":len(placeholders)},ensure_ascii=False)); return 2 if placeholders else 0

if __name__=="__main__": raise SystemExit(main())
