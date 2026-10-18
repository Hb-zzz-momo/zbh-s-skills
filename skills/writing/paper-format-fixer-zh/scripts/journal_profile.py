"""Semantic DOCX profiles, effective-format audit, and content-preserving repair.

python-docx is used for supported Word parts. Unknown complex parts fail closed
in profile mode; legacy conservative mode remains available separately.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
TEXT_PARTS = ('/word/document.xml', '/word/header', '/word/footer')


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def validate_profile(profile):
    if profile.get('schema_version') != 1:
        raise ValueError('Unsupported profile schema_version')
    for key in ('profile_id', 'label', 'authority', 'styles', 'page', 'required_styles', 'tables', 'render'):
        if not profile.get(key): raise ValueError('Profile missing ' + key)
    if profile['authority'] not in ('generic-demonstration', 'journal', 'school', 'course', 'user'):
        raise ValueError('Unknown format authority')
    if profile['authority'] != 'generic-demonstration' and not profile.get('sources'):
        raise ValueError('An external format authority requires source records')
    allowed_style = {'font_name', 'east_asia_font', 'font_size_pt', 'font_color', 'bold', 'italic',
        'alignment', 'line_spacing', 'line_spacing_pt', 'space_before_pt', 'space_after_pt',
        'first_line_indent_cm', 'left_indent_cm', 'right_indent_cm', 'keep_with_next',
        'keep_together', 'widow_control', 'character_spacing_pt', 'character_scale', 'kind'}
    for name, rule in profile['styles'].items():
        extra = set(rule) - allowed_style
        if extra: raise ValueError(f'Unsupported style fields in {name}: {sorted(extra)}')
        if not rule.get('font_name') or not rule.get('east_asia_font'):
            raise ValueError(f'Both Latin and East Asian fonts required in {name}')
        if not 6 <= rule.get('font_size_pt', 0) <= 72:
            raise ValueError(f'Invalid font size in {name}')
        if ('line_spacing' in rule) == ('line_spacing_pt' in rule):
            raise ValueError(f'Specify one line spacing unit in {name}')
        if rule.get('line_spacing', rule.get('line_spacing_pt', 0)) <= 0:
            raise ValueError(f'Non-positive line spacing in {name}')
        if rule.get('alignment') not in ALIGN: raise ValueError('Invalid alignment')
        if not re.fullmatch('[0-9A-Fa-f]{6}', rule.get('font_color', '')):
            raise ValueError('RGB color must contain six hex digits')
    if set(profile['required_styles']) - set(profile['styles']):
        raise ValueError('Required styles are not defined')
    for key in ('width_cm', 'height_cm', 'top_cm', 'bottom_cm', 'left_cm', 'right_cm'):
        if profile['page'].get(key, 0) <= 0: raise ValueError('Invalid page ' + key)
    if profile['page']['width_cm'] <= profile['page']['left_cm'] + profile['page']['right_cm']:
        raise ValueError('Page margins leave no usable width')
    if profile['page']['height_cm'] <= profile['page']['top_cm'] + profile['page']['bottom_cm']:
        raise ValueError('Page margins leave no usable height')
    if profile['tables'].get('kind') not in ('three_line', 'grid'):
        raise ValueError('Unsupported table kind')
    if not 0 < profile['render'].get('minimum_text_coverage', 0) <= 1:
        raise ValueError('Invalid text coverage threshold')
    return profile


def load_profile(path):
    return validate_profile(json.loads(Path(path).read_text(encoding='utf-8-sig')))


def story_parts(doc):
    for part in doc.part.package.parts:
        name = str(part.partname)
        if name == TEXT_PARTS[0] or name.startswith(TEXT_PARTS[1:]):
            if hasattr(part, 'element'): yield name, part.element


def iter_paragraphs(doc):
    for partname, element in story_parts(doc):
        for p in element.xpath('.//w:p'):
            yield partname, Paragraph(p, doc)


def iter_runs(paragraph):
    # Includes hyperlink, field, and content-control descendants.
    for element in paragraph._p.xpath('.//w:r'):
        yield Run(element, paragraph)


def plain_text(paragraph):
    return ''.join(t.text or '' for t in paragraph._p.xpath('.//w:t'))


def has_payload(paragraph):
    return bool(plain_text(paragraph).strip() or paragraph._p.xpath(
        './/w:drawing|.//w:pict|.//w:object|.//w:fldChar|.//w:br|.//w:sectPr|.//m:oMath'))


def style_chain(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id); yield style
        style = style.base_style


def effective_run_value(doc, paragraph, run, tag, attr='val', default=None):
    roots = [run._r.rPr]
    char = run._r.xpath('./w:rPr/w:rStyle/@w:val')
    if char:
        style = next((s for s in doc.styles if s.style_id == char[0]), None)
        if style is not None: roots += [s.element.rPr for s in style_chain(style)]
    roots += [s.element.rPr for s in style_chain(paragraph.style)]
    roots += doc.styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr')
    for root in roots:
        if root is None: continue
        node = root.find(qn('w:' + tag))
        if node is not None:
            value = node.get(qn('w:' + attr))
            if value is not None: return value
            if tag in ('b','i','bCs','iCs'):
                return '1'  # An empty on/off element means enabled in OOXML.
            if tag == 'rFonts' and node.get(qn('w:' + attr + 'Theme')):
                return 'THEME:' + node.get(qn('w:' + attr + 'Theme'))
            if tag == 'color' and node.get(qn('w:themeColor')):
                return 'THEME:' + node.get(qn('w:themeColor'))
    return default


def effective_format(paragraph, attr):
    value = getattr(paragraph.paragraph_format, attr)
    if value is not None: return value
    for style in style_chain(paragraph.style):
        value = getattr(style.paragraph_format, attr)
        if value is not None: return value
    return None


def set_child(parent, tag, attrs):
    node = parent.find(qn('w:' + tag))
    if node is None: node = OxmlElement('w:' + tag); parent.append(node)
    node.attrib.clear()
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    return node


def set_font(element, rule):
    pr = element.get_or_add_rPr()
    set_child(pr, 'rFonts', {'ascii': rule['font_name'], 'hAnsi': rule['font_name'],
                           'cs': rule['font_name'], 'eastAsia': rule['east_asia_font']})
    set_child(pr, 'sz', {'val': round(rule['font_size_pt'] * 2)})
    set_child(pr, 'szCs', {'val': round(rule['font_size_pt'] * 2)})
    set_child(pr, 'color', {'val': rule['font_color']})
    set_child(pr, 'spacing', {'val': round(rule.get('character_spacing_pt', 0) * 20)})
    set_child(pr, 'w', {'val': rule.get('character_scale', 100)})


def set_paragraph_format(fmt, rule):
    fmt.alignment = ALIGN[rule['alignment']]
    fmt.line_spacing = Pt(rule['line_spacing_pt']) if 'line_spacing_pt' in rule else rule['line_spacing']
    if 'line_spacing_pt' in rule: fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for field in ('space_before', 'space_after'):
        setattr(fmt, field, Pt(rule.get(field + '_pt', 0)))
    for field in ('first_line_indent', 'left_indent', 'right_indent'):
        setattr(fmt, field, Cm(rule.get(field + '_cm', 0)))
    fmt.keep_with_next = rule.get('keep_with_next', False)
    fmt.keep_together = rule.get('keep_together', False)
    fmt.widow_control = rule.get('widow_control', True)


def prepare_styles(doc, profile):
    for name, rule in profile['styles'].items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != 'Normal': style.base_style = doc.styles['Normal']
        set_font(style.element, rule)
        style.font.bold = rule.get('bold', False)
        style.font.italic = rule.get('italic', False)
        set_paragraph_format(style.paragraph_format, rule)
        pr = style.element.get_or_add_pPr()
        for tag in ('pBdr', 'shd'):  # A generic profile does not inherit decorative rules.
            for node in pr.findall(qn('w:' + tag)): pr.remove(node)


def all_tables(doc):
    for partname, element in story_parts(doc):
        for item in element.xpath('.//w:tbl'): yield partname, Table(item, doc)


def apply_table(table, config, usable_cm):
    if config.get('kind') not in ('three_line', 'grid'): raise ValueError('Invalid table override')
    widths = config.get('widths_cm')
    if not widths:
        existing = [c.width.cm if c.width else 0 for c in table.columns]
        widths = existing if all(w > 0 for w in existing) else [usable_cm / len(table.columns)] * len(table.columns)
        if sum(widths) > usable_cm:
            widths = [w * usable_cm / sum(widths) for w in widths]
    if len(widths) != len(table.columns) or any(w <= 0 for w in widths) or sum(widths) > usable_cm + 0.02:
        raise ValueError('Table column widths do not fit the page')
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Normal Table'
    tblpr = table._tbl.tblPr
    set_child(tblpr, 'tblW', {'type': 'dxa', 'w': round(sum(widths) * 567)})
    borders = set_child(tblpr, 'tblBorders', {})
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        set_child(borders, edge, {'val': 'nil'})
    margins = set_child(tblpr, 'tblCellMar', {})
    for edge in ('top', 'bottom', 'left', 'right'):
        set_child(margins, edge, {'type': 'dxa', 'w': 70 if edge in ('top', 'bottom') else 90})
    for col, width in zip(table.columns, widths): col.width = Cm(width)
    seen = set()
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        set_child(trpr, 'cantSplit', {'val': 1})
        for item in trpr.findall(qn('w:trHeight')): trpr.remove(item)
        for item in trpr.findall(qn('w:tblHeader')): trpr.remove(item)
        if i == 0: set_child(trpr, 'tblHeader', {'val': 1})
        for j, cell in enumerate(row.cells):
            if cell._tc in seen: continue
            seen.add(cell._tc)
            cell.width = Cm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            set_child(pr, 'shd', {'val': 'clear', 'color': 'auto', 'fill': 'FFFFFF'})
            cb = set_child(pr, 'tcBorders', {})
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV', 'start', 'end'):
                visible = config['kind'] == 'grid' or (edge == 'top' and i == 0) or (
                    edge == 'bottom' and i in (0, len(table.rows)-1))
                set_child(cb, edge, {'val': 'single' if visible else 'nil',
                    'sz': 8 if i in (0, len(table.rows)-1) else 4, 'color': '000000'})
            for p in cell.paragraphs: p.style = 'Table Text'


def table_details(table):
    visible = lambda n: n is not None and n.get(qn('w:val'), 'nil') not in ('nil', 'none', '0')
    three_line = True; vertical = False; rules = []
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV', 'start', 'end'):
                node = cell._tc.find('./' + qn('w:tcPr') + '/' + qn('w:tcBorders') + '/' + qn('w:' + edge))
                if node is None: node = table._tbl.tblPr.find('./' + qn('w:tblBorders') + '/' + qn('w:' + edge))
                on = visible(node)
                expected = (edge == 'top' and i == 0) or (edge == 'bottom' and i in (0, len(table.rows)-1))
                if on != expected: three_line = False
                if on and edge in ('left', 'right', 'insideV', 'start', 'end'): vertical = True
                if on: rules.append({'row': i+1, 'col': j+1, 'edge': edge})
    repeat = table.rows[0]._tr.find('./' + qn('w:trPr') + '/' + qn('w:tblHeader'))
    return {'rows': len(table.rows), 'columns': len(table.columns), 'three_line': three_line,
        'vertical_rules': vertical, 'visible_rules': rules,
        'repeat_header': repeat is not None and repeat.get(qn('w:val'), '1') not in ('0', 'false'),
        'cant_split_rows': sum(bool(row._tr.xpath('./w:trPr/w:cantSplit[not(@w:val="0")]')) for row in table.rows),
        'max_cell_characters': max((len(c.text) for row in table.rows for c in row.cells), default=0),
        'width_cm': round(sum(c.width.cm for c in table.columns if c.width), 3)}


def table_keep_map(doc, profile, semantic_map=None):
    result = {}
    for index, (_, table) in enumerate(all_tables(doc), 1):
        config = dict(profile['tables']); config.update((semantic_map or {}).get('tables', {}).get(str(index), {}))
        short = len(table.rows) <= config.get('keep_together_max_rows', 8) and max(
            (len(c.text) for row in table.rows for c in row.cells), default=0) <= config.get('max_cell_characters',280)
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs: result[p._p] = short and index < len(table.rows)-1
    return result


def context_rule(p, profile, table_keep):
    rule = dict(profile['styles'][p.style.name])
    if p._p in table_keep: rule['keep_with_next'] = table_keep[p._p]
    if p.style.name == 'Caption':
        nxt = p._p.getnext()
        rule['keep_with_next'] = nxt is not None and (nxt.tag == qn('w:tbl') or (
            nxt.tag == qn('w:p') and Paragraph(nxt,p._parent).style.name == 'SourceNote'))
    return rule


def content_signature(doc):
    # Formatting must leave text, field instructions, media, and link targets intact.
    records = []
    for name, element in story_parts(doc):
        records.append((name, element.xpath('.//w:t/text()'), element.xpath('.//w:instrText/text()'),
                        len(element.xpath('.//w:drawing|.//w:pict|.//w:object'))))
    for part in doc.part.package.parts:
        if str(part.partname).startswith('/word/media/'):
            records.append((str(part.partname), hashlib.sha256(part.blob).hexdigest()))
        for rel in part.rels.values():
            if rel.is_external: records.append((str(part.partname), rel.reltype, rel.target_ref))
    return hashlib.sha256(json.dumps(records, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def unsupported_features(doc):
    result = []
    for part in doc.part.package.parts:
        name = str(part.partname)
        if name.startswith(('/word/footnotes', '/word/endnotes', '/word/comments')):
            result.append(name)
    for name, element in story_parts(doc):
        for query, reason in [('.//w:ins|.//w:del', 'tracked changes'),
                              ('.//w:txbxContent', 'text boxes'), ('.//m:oMath', 'native equations'),
                              ('.//w:altChunk', 'imported content')]:
            if element.xpath(query): result.append(name + ': ' + reason)
    return result


def apply_profile(doc, profile, semantic_map=None):
    validate_profile(profile)
    unsupported = unsupported_features(doc)
    if unsupported: raise ValueError('Profile repair requires a specialized path: ' + '; '.join(unsupported))
    mapping = semantic_map or {}
    before = content_signature(doc)
    prepare_styles(doc, profile)
    for index, name in mapping.get('paragraphs', {}).items():
        if name not in profile['styles']: raise ValueError('Unknown semantic style: ' + name)
        doc.paragraphs[int(index)-1].style = name
    page_keys = {'width_cm': 'page_width', 'height_cm': 'page_height', 'top_cm': 'top_margin',
                'bottom_cm': 'bottom_margin', 'left_cm': 'left_margin', 'right_cm': 'right_margin',
                'header_cm': 'header_distance', 'footer_cm': 'footer_distance'}
    for section in doc.sections:
        for key, attr in page_keys.items():
            if key in profile['page']: setattr(section, attr, Cm(profile['page'][key]))
    usable = profile['page']['width_cm'] - profile['page']['left_cm'] - profile['page']['right_cm']
    for index, (_, table) in enumerate(all_tables(doc), 1):
        config = dict(profile['tables']); config.update(mapping.get('tables', {}).get(str(index), {}))
        apply_table(table, config, usable)
    table_keep = table_keep_map(doc, profile, mapping)
    for name, p in iter_paragraphs(doc):
        style = p.style.name
        if name.startswith(('/word/header', '/word/footer')): style = 'Footer'; p.style = style
        if style not in profile['styles']:
            if has_payload(p): raise ValueError('Unmapped semantic style ' + style)
            continue
        rule = context_rule(p, profile, table_keep)
        set_paragraph_format(p.paragraph_format, rule)
        pr = p._p.get_or_add_pPr()
        for tag in ('pBdr', 'shd', 'contextualSpacing', 'snapToGrid'):
            for item in pr.findall(qn('w:' + tag)): pr.remove(item)
        for run in iter_runs(p):
            set_font(run._r, rule)
            # Preserve deliberate inline emphasis in body and references.
            if rule.get('kind') in ('title', 'heading', 'caption', 'keywords'):
                run.bold = rule.get('bold', False); run.italic = rule.get('italic', False)
    if content_signature(doc) != before: raise ValueError('Formatting changed content or media')
    return doc


def audit_profile(doc, profile, semantic_map=None):
    validate_profile(profile)
    errors = []; warnings = []; used = Counter(); paragraphs = []; fonts = Counter()
    def issue(code, where, expected=None, actual=None):
        errors.append({'code': code, 'location': where, 'expected': expected, 'actual': actual})
    for feature in unsupported_features(doc): issue('unsupported_feature', feature)
    table_keep = table_keep_map(doc, profile, semantic_map)
    for i, (part, p) in enumerate(iter_paragraphs(doc), 1):
        if not has_payload(p): continue
        style = p.style.name; used[style] += 1
        where = f'{part}:p{i}'
        if style not in profile['styles']:
            issue('unmapped_style', where, list(profile['styles']), style); continue
        rule = context_rule(p, profile, table_keep)
        paragraphs.append({'location': where, 'style': style, 'text': plain_text(p)[:120]})
        for run in iter_runs(p):
            if not run.text.strip() and not run._r.xpath('.//w:instrText'): continue
            for code, tag, attr, expected, factor in [
                ('font_name', 'rFonts', 'ascii', rule['font_name'], None),
                ('font_size', 'sz', 'val', rule['font_size_pt'], 0.5),
                ('font_color', 'color', 'val', rule['font_color'], None),
                ('character_spacing', 'spacing', 'val', rule.get('character_spacing_pt', 0), 0.05),
                ('character_scale', 'w', 'val', rule.get('character_scale', 100), 1)]:
                fallback = {'character_spacing': '0', 'character_scale': '100'}.get(code)
                actual = effective_run_value(doc, p, run, tag, attr, fallback)
                if factor is not None and actual is not None: actual = float(actual) * factor
                same = abs(actual - expected) < 0.02 if isinstance(actual, (float,int)) else actual == expected
                if not same: issue(code, where, expected, actual)
            font = effective_run_value(doc, p, run, 'rFonts', 'ascii')
            fonts[str(font)] += 1
            if rule.get('kind') in ('title','heading','caption','keywords'):
                for key,tag in [('bold','b'),('italic','i')]:
                    raw=effective_run_value(doc,p,run,tag,default='0')
                    actual=raw not in ('0','false','off')
                    if actual!=rule.get(key,False):issue('font_'+key,where,rule.get(key,False),actual)
            if re.search('[\u3400-\u9fff]', run.text):
                east = effective_run_value(doc, p, run, 'rFonts', 'eastAsia')
                if east != rule['east_asia_font']: issue('east_asia_font', where, rule['east_asia_font'], east)
        for attr in ('alignment', 'line_spacing', 'space_before', 'space_after', 'first_line_indent',
                     'left_indent', 'right_indent', 'keep_with_next', 'keep_together', 'widow_control'):
            value = effective_format(p, attr)
            if attr == 'alignment': expected = int(ALIGN[rule['alignment']]); actual = int(value) if value is not None else 0
            elif attr == 'line_spacing':
                expected = rule.get('line_spacing_pt', rule.get('line_spacing'))
                actual = value.pt if hasattr(value, 'pt') else value
                fixed = hasattr(value, 'pt')
                if fixed != ('line_spacing_pt' in rule): issue('line_spacing_unit', where, 'pt' if 'line_spacing_pt' in rule else 'multiple', 'pt' if fixed else 'multiple')
            elif attr in ('space_before', 'space_after'):
                expected = rule.get(attr + '_pt', 0); actual = value.pt if value is not None else 0
            elif attr.endswith('indent'):
                expected = rule.get(attr + '_cm', 0); actual = value.cm if value is not None else 0
            else: expected = rule.get(attr, attr == 'widow_control'); actual = bool(value) if value is not None else attr == 'widow_control'
            if actual is None or abs(float(actual) - float(expected)) > 0.025:
                issue(attr, where, expected, actual)
        if 'line_spacing_pt' in rule and effective_format(p,'line_spacing_rule') != WD_LINE_SPACING.EXACTLY:
            issue('line_spacing_rule',where,'EXACTLY',str(effective_format(p,'line_spacing_rule')))
    for name in profile['required_styles']:
        if not used[name]: issue('required_style', 'document', name, 'absent')
    page_mapping = {'width_cm':'page_width','height_cm':'page_height','top_cm':'top_margin',
                    'bottom_cm':'bottom_margin','left_cm':'left_margin','right_cm':'right_margin'}
    for index, section in enumerate(doc.sections, 1):
        for key, attr in page_mapping.items():
            actual = getattr(section, attr).cm
            if abs(actual - profile['page'][key]) > .025: issue('page_' + key, f'section{index}', profile['page'][key], actual)
        cols = section._sectPr.find(qn('w:cols'))
        if cols is not None and cols.get(qn('w:num'), '1') != '1': issue('page_columns', f'section{index}', 1, cols.get(qn('w:num')))
    tables = []; usable = profile['page']['width_cm'] - profile['page']['left_cm'] - profile['page']['right_cm']
    for index, (_, table) in enumerate(all_tables(doc), 1):
        details = table_details(table); details['index'] = index; tables.append(details)
        config = dict(profile['tables']); config.update((semantic_map or {}).get('tables', {}).get(str(index), {}))
        if config['kind'] == 'three_line' and not details['three_line']: issue('table_borders', f'table{index}', 'three_line', details['visible_rules'])
        if not details['repeat_header']: issue('table_repeat_header', f'table{index}', True, False)
        if details['cant_split_rows'] != details['rows']: issue('table_row_split', f'table{index}')
        if details['width_cm'] > usable + .03: issue('table_width', f'table{index}', usable, details['width_cm'])
        if details['max_cell_characters'] > config.get('max_cell_characters', 280):
            warnings.append({'code':'table_density','location':f'table{index}','actual':details['max_cell_characters'],
                             'action':'Review the rendered table; move explanatory prose to the body if cramped.'})
        prev = table._tbl.getprevious()
        if prev is None or prev.tag != qn('w:p') or Paragraph(prev,doc).style.name != 'Caption':
            issue('table_caption_position', f'table{index}', 'Caption immediately above table')
    figures = doc.element.xpath('.//w:p[w:r/w:drawing]')
    for i, element in enumerate(figures, 1):
        nxt = element.getnext()
        if nxt is None or nxt.tag != qn('w:p') or Paragraph(nxt,doc).style.name != 'Caption':
            issue('figure_caption_position', f'figure{i}', 'Caption immediately below figure')
    for i, shape in enumerate(doc.inline_shapes, 1):
        if shape.width.cm > usable + .03: issue('figure_width', f'figure{i}', usable, shape.width.cm)
    return {'status':'BLOCK' if errors else ('CAUTION' if warnings else 'PASS'),
            'profile_id':profile['profile_id'], 'authority':profile['authority'],
            'scope':'DOCX structure and effective formatting; visual inspection remains required',
            'errors':errors,'warnings':warnings,'used_styles':dict(used),'effective_latin_fonts':dict(fonts),
            'paragraphs':paragraphs,'tables':tables,'figure_count':len(doc.inline_shapes)}


def format_file(input_path, output_path, profile, semantic_map=None, overwrite=False):
    input_path = Path(input_path); output_path = Path(output_path)
    if input_path.resolve() == output_path.resolve(): raise ValueError('Use a separate output path; in-place profile repair is disabled')
    if output_path.exists() and not overwrite: raise FileExistsError(output_path)
    doc = Document(input_path); signature = content_signature(doc)
    apply_profile(doc, profile, semantic_map)
    output_path.parent.mkdir(parents=True, exist_ok=True); doc.save(output_path)
    final = Document(output_path)
    if content_signature(final) != signature: raise ValueError('Saved output changed content or media')
    return {'input':str(input_path.resolve()),'output':str(output_path.resolve()),
            'input_sha256':sha256(input_path),'output_sha256':sha256(output_path),
            'content_preserved':True,'profile_id':profile['profile_id'],
            'conformance':audit_profile(final, profile, semantic_map)}
