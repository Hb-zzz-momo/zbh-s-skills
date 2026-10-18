"""Refuse release claims when structural, render, or visual evidence is stale."""
import argparse
import json
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from journal_profile import audit_profile, load_profile, sha256, write_json


def validate_visual_review(render, review):
    reasons=[]
    if review.get('docx_sha256') != render.get('docx_sha256'): reasons.append('visual_docx_stale')
    if review.get('pdf_sha256') != render.get('pdf_sha256'): reasons.append('visual_pdf_stale')
    if not review.get('reviewer'): reasons.append('visual_reviewer_missing')
    expected={p['number']:p['sha256'] for p in render.get('pages',[])}
    entries=review.get('pages',[])
    if not expected or len(entries)!=len(expected) or {p.get('number') for p in entries}!=set(expected):
        reasons.append('visual_pages_incomplete')
    for entry in entries:
        if entry.get('sha256')!=expected.get(entry.get('number')): reasons.append('visual_image_stale')
        if entry.get('status')!='PASS' or not entry.get('notes'): reasons.append('visual_page_unaccepted')
    return sorted(set(reasons))


def verify_release(docx, profile_path, audit_path, render_path, visual_path, semantic_map=None):
    profile=load_profile(profile_path); current=sha256(docx); profile_hash=sha256(profile_path)
    actual=audit_profile(Document(docx),profile,semantic_map)
    reasons=[]
    if actual['errors']: reasons.append('structure_failed')
    def read(path):
        return json.loads(Path(path).read_text(encoding='utf-8-sig')) if path and Path(path).exists() else None
    audit=read(audit_path); render=read(render_path); visual=read(visual_path)
    if audit is None: reasons.append('audit_missing')
    elif audit.get('docx_sha256')!=current or audit.get('profile_sha256')!=profile_hash: reasons.append('audit_stale')
    if render is None: reasons.append('render_missing')
    else:
        if render.get('docx_sha256')!=current or render.get('profile_sha256')!=profile_hash: reasons.append('render_stale')
        if render.get('status')!='PASS': reasons.append('render_failed')
        pdf=Path(render.get('pdf',''))
        if not pdf.is_file() or sha256(pdf)!=render.get('pdf_sha256'): reasons.append('render_pdf_stale')
        else:
            if len(PdfReader(pdf).pages)!=render.get('page_count'): reasons.append('render_page_count')
            # Check the actual PDF again, rather than trusting editable receipt fields.
            import pdfplumber
            from journal_profile import story_parts
            from record_render import text_coverage
            source_doc=Document(docx)
            source_text='\n'.join(' '.join(e.xpath('.//w:t/text()')) for _,e in story_parts(source_doc))
            with pdfplumber.open(pdf) as actual_pdf:
                rendered_text='\n'.join(p.extract_text() or '' for p in actual_pdf.pages)
                font_names={c['fontname'].split('+')[-1] for p in actual_pdf.pages for c in p.chars}
            if text_coverage(source_text,rendered_text)<profile['render']['minimum_text_coverage']:
                reasons.append('render_text_coverage')
            if any(not any(font.startswith(prefix) for prefix in profile['render']['allowed_pdf_font_prefixes'])
                   for font in font_names):reasons.append('render_font_fallback')
        pages=render.get('pages',[])
        if [p.get('number') for p in pages]!=list(range(1,render.get('page_count',0)+1)): reasons.append('render_images_incomplete')
        for p in pages:
            file=Path(p.get('path',''))
            if not file.is_file() or sha256(file)!=p.get('sha256'): reasons.append('render_image_stale')
        if not render.get('read_only_copy'): reasons.append('render_copy_unverified')
        if render.get('text_coverage',0)<profile['render']['minimum_text_coverage']: reasons.append('render_text_coverage')
        if render.get('unknown_fonts'): reasons.append('render_font_fallback')
    if visual is None: reasons.append('visual_missing')
    elif render is not None:
        reasons.extend(validate_visual_review(render,visual))
        resolved=set(visual.get('resolved_warnings',[]))
        unresolved={w['code']+':'+w['location'] for w in actual['warnings']} - resolved
        if unresolved: reasons.append('visual_warnings_unresolved')
    if profile['authority']=='generic-demonstration':
        from journal_profile import plain_text
        notice=' '.join(plain_text(p) for p in Document(docx).paragraphs if p.style.name=='Subtitle').lower()
        if not ('通用' in notice or 'generic' in notice): reasons.append('generic_label_missing')
    return {'status':'BLOCK' if reasons else 'PASS','reasons':sorted(set(reasons)),
        'docx_sha256':current,'profile_sha256':profile_hash,'profile_id':profile['profile_id'],
        'authority':profile['authority'],'structure':actual['status'],
        'scope':'Format validated against this specific profile and these rendered pages. This is not acceptance by a journal or a research contribution assessment.'}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    for name in ('docx','profile','out'): parser.add_argument('--'+name,required=True)
    for name in ('audit','render','visual','map'): parser.add_argument('--'+name)
    args=parser.parse_args()
    mapping=json.loads(Path(args.map).read_text(encoding='utf-8')) if args.map else None
    result=verify_release(args.docx,args.profile,args.audit,args.render,args.visual,mapping)
    write_json(args.out,result)
    return 1 if result['status']=='BLOCK' else 0


if __name__=='__main__': raise SystemExit(main())
