"""Build an editable manuscript from reviewed semantic blocks, not free-form HTML."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from journal_profile import prepare_styles, apply_profile, audit_profile, load_profile, write_json, sha256


def add_paragraph(doc, block, profile):
    style=block.get('style','Normal')
    if style not in profile['styles']: raise ValueError('Unknown paragraph style ' + style)
    p=doc.add_paragraph(style=style)
    runs=block.get('runs', [{'text':block.get('text','')}])
    for item in runs:
        r=p.add_run(item['text'])
        if item.get('bold') is not None: r.bold=item['bold']
        if item.get('italic') is not None: r.italic=item['italic']
        if item.get('url'):
            h=OxmlElement('w:hyperlink')
            h.set(qn('r:id'),doc.part.relate_to(item['url'],RT.HYPERLINK,is_external=True))
            p._p.remove(r._r);h.append(r._r);p._p.append(h)
    return p


def build(data_path, profile_path, out_path, overwrite=False):
    data_path=Path(data_path);out_path=Path(out_path);profile=load_profile(profile_path)
    if out_path.exists() and not overwrite: raise FileExistsError(out_path)
    if data_path.resolve()==out_path.resolve(): raise ValueError('Input and output must differ')
    data=json.loads(data_path.read_text(encoding='utf-8-sig'))
    if not isinstance(data.get('blocks'),list) or not data['blocks']: raise ValueError('Manuscript blocks required')
    doc=Document();prepare_styles(doc,profile)
    doc.core_properties.author=data.get('author','')
    doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    doc.core_properties.title=next((b.get('text','') for b in data['blocks'] if b.get('style')=='Title'),'')
    mapping={'tables':{}}
    usable=profile['page']['width_cm']-profile['page']['left_cm']-profile['page']['right_cm']
    for block in data['blocks']:
        kind=block.get('type','paragraph')
        if kind=='paragraph': add_paragraph(doc,block,profile)
        elif kind=='table':
            rows=block['rows']
            if not rows or len(rows)<2 or not rows[0] or any(len(r)!=len(rows[0]) for r in rows):
                raise ValueError('A table requires a header and equally sized rows')
            add_paragraph(doc,{'text':block['caption'],'style':'Caption'},profile)
            table=doc.add_table(rows=len(rows),cols=len(rows[0]))
            for i,row in enumerate(rows):
                for j,value in enumerate(row): table.cell(i,j).text=str(value)
            index=str(len(doc.tables));mapping['tables'][index]={}
            for key in ('kind','widths_cm'):
                if key in block: mapping['tables'][index][key]=block[key]
            if block.get('note'): add_paragraph(doc,{'text':block['note'],'style':'SourceNote'},profile)
        elif kind=='figure':
            image=Path(block['image'])
            if not image.is_absolute(): image=data_path.parent/image
            width=block.get('width_cm',usable)
            if not 0<width<=usable: raise ValueError('Figure width must fit the page')
            p=doc.add_paragraph(style='Figure');p.add_run().add_picture(str(image),width=Cm(width))
            add_paragraph(doc,{'text':block['caption'],'style':'Caption'},profile)
            if block.get('note'): add_paragraph(doc,{'text':block['note'],'style':'SourceNote'},profile)
        else: raise ValueError('Unsupported manuscript block type ' + str(kind))
    footer=doc.sections[0].footer.paragraphs[0];footer.style='Footer'
    r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    result=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';result.append(text);field.append(result)
    footer._p.append(field)
    apply_profile(doc,profile,mapping)
    for table in doc.tables:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:r.bold=True
    out_path.parent.mkdir(parents=True,exist_ok=True);doc.save(out_path)
    result={'docx':str(out_path.resolve()),'docx_sha256':sha256(out_path),
            'profile_sha256':sha256(profile_path),'source_json_sha256':sha256(data_path),
            'semantic_map':mapping,'conformance':audit_profile(Document(out_path),profile,mapping)}
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for name in ('data','profile','out','report'):p.add_argument('--'+name,required=True)
    p.add_argument('--overwrite',action='store_true');a=p.parse_args()
    result=build(a.data,a.profile,a.out,a.overwrite);write_json(a.report,result)
    return 1 if result['conformance']['errors'] else 0


if __name__=='__main__':raise SystemExit(main())
