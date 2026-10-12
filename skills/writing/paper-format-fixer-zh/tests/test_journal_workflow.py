"""Regression tests for observable document and release invariants."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import audit_docx_format as old_audit
import fix_docx_format as old_fix


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def profile(self, language='en'):
        from journal_profile import load_profile
        return load_profile(ROOT / 'references' / ('general-journal-' + language + '.json'))

    def sample(self, language='en'):
        from journal_profile import prepare_styles, apply_profile
        doc = Document()
        profile = self.profile(language)
        prepare_styles(doc, profile)
        for text, style in [('BYD profitability and expansion', 'Title'),
                            ('Abstract', 'Abstract Heading'),
                            ('A descriptive analysis using public company disclosures.', 'Abstract'),
                            ('Keywords: BYD; profitability; case study', 'Keywords'),
                            ('1 Introduction', 'Heading 1'),
                            ('Revenue grew while the profit margin declined.', 'Normal'),
                            ('References', 'References Heading'),
                            ('BYD (2026). Annual report 2025.', 'Reference')]:
            doc.add_paragraph(text, style)
        apply_profile(doc, profile)
        return doc, profile

    def test_fixed_line_spacing_uses_points(self):
        self.assertEqual(old_audit.line_spacing_value(Pt(24)), '24.0 pt')

    def test_image_only_paragraph_is_not_deleted(self):
        doc = Document()
        doc.add_paragraph('')
        path = self.work / 'image.png'
        Image.new('RGB', (20, 20), 'white').save(path)
        doc.add_paragraph().add_run().add_picture(str(path))
        old_fix.collapse_blank_paragraphs(doc)
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_hyperlink_color_is_audited_and_fixed(self):
        doc = Document()
        p = doc.add_paragraph('Reference ')
        h = OxmlElement('w:hyperlink')
        r = OxmlElement('w:r')
        pr = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '0000EE')
        pr.append(color); r.append(pr)
        t = OxmlElement('w:t'); t.text = 'source'; r.append(t); h.append(r); p._p.append(h)
        self.assertIn('0000EE', old_audit.collect_run_colors(doc))
        old_fix.force_font_color(doc)
        self.assertNotIn('0000EE', old_audit.collect_run_colors(doc))

    def test_inherited_fonts_are_resolved(self):
        from journal_profile import audit_profile
        doc, profile = self.sample()
        run = doc.paragraphs[5].runs[0]
        run._r.remove(run._r.rPr)
        self.assertEqual(audit_profile(doc, profile)['errors'], [])

    def test_wrong_character_spacing_fails(self):
        from journal_profile import audit_profile
        doc, profile = self.sample()
        item = doc.paragraphs[5].runs[0]._r.get_or_add_rPr().find(qn('w:spacing'))
        item.set(qn('w:val'), '40')
        codes = [x['code'] for x in audit_profile(doc, profile)['errors']]
        self.assertIn('character_spacing', codes)

    def test_wrong_character_scale_fails(self):
        from journal_profile import audit_profile
        doc, profile = self.sample()
        item = doc.paragraphs[5].runs[0]._r.get_or_add_rPr().find(qn('w:w'))
        item.set(qn('w:val'), '80')
        self.assertIn('character_scale', [x['code'] for x in audit_profile(doc, profile)['errors']])

    def test_chinese_east_asian_font_fails_when_wrong(self):
        from journal_profile import audit_profile
        doc, profile = self.sample('zh')
        run = doc.paragraphs[5].runs[0]; run.text = '比亚迪利润率下降'
        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'MissingCJKFont')
        self.assertIn('east_asia_font', [x['code'] for x in audit_profile(doc, profile)['errors']])

    def test_three_lines_are_real_cell_borders(self):
        from journal_profile import apply_profile, audit_profile, table_details
        doc, profile = self.sample()
        table = doc.add_table(rows=3, cols=2)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells): cell.text = str(i * 2 + j)
        doc.add_paragraph('Table 1 Financial indicators', 'Caption')
        # The caption must precede this table.
        table._tbl.addprevious(doc.paragraphs[-1]._p)
        apply_profile(doc, profile)
        info = table_details(table)
        self.assertTrue(info['three_line'])
        self.assertTrue(info['repeat_header'])
        self.assertEqual(audit_profile(doc, profile)['errors'], [])

    def test_vertical_table_border_fails(self):
        from journal_profile import apply_profile, audit_profile
        doc, profile = self.sample()
        p = doc.add_paragraph('Table 1 Financial indicators', 'Caption')
        table = doc.add_table(rows=2, cols=2)
        apply_profile(doc, profile)
        borders = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
        borders.find(qn('w:left')).set(qn('w:val'), 'single')
        self.assertIn('table_borders', [x['code'] for x in audit_profile(doc, profile)['errors']])

    def test_large_table_cell_triggers_density_warning(self):
        from journal_profile import apply_profile, audit_profile
        doc, profile = self.sample()
        doc.add_paragraph('Table 1 Long prose', 'Caption')
        table = doc.add_table(rows=2, cols=2)
        table.cell(1, 0).text = 'a' * 600
        apply_profile(doc, profile)
        self.assertIn('table_density', [x['code'] for x in audit_profile(doc, profile)['warnings']])

    def test_profile_formatting_preserves_text_and_media(self):
        from journal_profile import apply_profile, content_signature
        doc, profile = self.sample()
        path = self.work / 'image.png'; Image.new('RGB', (12, 12)).save(path)
        doc.add_paragraph(style='Figure').add_run().add_picture(str(path))
        before = content_signature(doc)
        apply_profile(doc, profile)
        self.assertEqual(before, content_signature(doc))

    def test_bad_profile_is_rejected(self):
        from journal_profile import validate_profile
        profile = self.profile(); profile['styles']['Normal']['line_spacing'] = 0
        with self.assertRaises(ValueError): validate_profile(profile)

    def test_missing_abstract_is_not_accepted(self):
        from journal_profile import audit_profile
        doc, profile = self.sample()
        doc.paragraphs[2].style = doc.styles['Normal']
        self.assertIn('required_style', [x['code'] for x in audit_profile(doc, profile)['errors']])

    def test_rendered_pdf_must_contain_manuscript_text(self):
        from record_render import text_coverage
        self.assertLess(text_coverage('BYD revenue and profitability declined', 'unrelated document'), 0.95)
        self.assertGreaterEqual(text_coverage('BYD revenue', 'BYD revenue page 1'), 0.95)

    def test_release_rejects_missing_and_stale_evidence(self):
        from verify_release import verify_release
        doc, profile = self.sample()
        path = self.work / 'sample.docx'; doc.save(path)
        profile_path = ROOT / 'references' / 'general-journal-en.json'
        result = verify_release(path, profile_path, None, None, None)
        self.assertEqual(result['status'], 'BLOCK')
        self.assertIn('render_missing', result['reasons'])

    def test_in_place_modification_is_rejected(self):
        from journal_profile import format_file
        doc, profile = self.sample(); path = self.work / 'same.docx'; doc.save(path)
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        with self.assertRaises(ValueError): format_file(path, path, profile, overwrite=True)
        self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_partial_page_review_is_blocked(self):
        from verify_release import validate_visual_review
        render = {'docx_sha256': 'd', 'pdf_sha256': 'p', 'pages': [
            {'number': 1, 'sha256': 'a'}, {'number': 2, 'sha256': 'b'}]}
        review = {'docx_sha256': 'd', 'pdf_sha256': 'p', 'reviewer': 'test',
                  'pages': [{'number': 1, 'sha256': 'a', 'status': 'PASS', 'notes': 'Read and checked.'}]}
        self.assertIn('visual_pages_incomplete', validate_visual_review(render, review))

    def test_old_visual_review_is_blocked(self):
        from verify_release import validate_visual_review
        render = {'docx_sha256': 'new', 'pdf_sha256': 'p', 'pages': []}
        review = {'docx_sha256': 'old', 'pdf_sha256': 'p', 'reviewer': 'test', 'pages': []}
        self.assertIn('visual_docx_stale', validate_visual_review(render, review))

    def test_builder_entry_point_creates_editable_tables(self):
        import subprocess
        data={'blocks':[
            {'style':'Title','text':'BYD profitability and expansion'},
            {'style':'Subtitle','text':'Generic journal manuscript demonstration'},
            {'style':'Abstract Heading','text':'Abstract'},
            {'style':'Abstract','text':'A descriptive analysis of public disclosures.'},
            {'style':'Keywords','text':'Keywords: BYD; profitability'},
            {'style':'Heading 1','text':'1 Introduction'},
            {'text':'Revenue and profitability are distinct indicators.'},
            {'type':'table','caption':'Table 1 Indicators','rows':[['Year','Revenue'],['2025','803.965']]},
            {'style':'References Heading','text':'References'},
            {'style':'Reference','text':'BYD (2026). Annual report 2025.'}]}
        input_path=self.work/'data.json';input_path.write_text(json.dumps(data),encoding='utf-8')
        output=self.work/'paper.docx';report=self.work/'build.json'
        result=subprocess.run([sys.executable,str(ROOT/'scripts'/'build_manuscript.py'),
            '--data',str(input_path),'--profile',str(ROOT/'references'/'general-journal-en.json'),
            '--out',str(output),'--report',str(report)],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        doc=Document(output)
        self.assertEqual(doc.core_properties.author,'')
        self.assertEqual(doc.tables[0].cell(1,1).text,'803.965')
        self.assertEqual(json.loads(report.read_text(encoding='utf-8'))['conformance']['errors'],[])

    def test_release_accepts_current_evidence_then_rejects_mutation(self):
        from journal_profile import write_json,sha256,audit_profile
        from verify_release import verify_release
        from record_render import record_render
        from reportlab.pdfgen.canvas import Canvas
        doc,profile=self.sample();doc.add_paragraph('Generic journal manuscript demonstration','Subtitle')
        path=self.work/'sample.docx';doc.save(path)
        prof=self.work/'profile.json'
        profile['render']['allowed_pdf_font_prefixes'].append('Times-Roman')
        write_json(prof,profile)
        pdf=self.work/'sample.pdf';canvas=Canvas(str(pdf));canvas.setFont('Times-Roman',12)
        for i,p in enumerate(doc.paragraphs):canvas.drawString(40,750-i*40,p.text)
        canvas.save()
        # Synthetic unit fixture; separate integration exercises real Word.
        png=self.work/'page-1.png';Image.new('RGB',(20,20)).save(png)
        audit=self.work/'audit.json';render=self.work/'render.json';visual=self.work/'visual.json'
        write_json(audit,{'docx_sha256':sha256(path),'profile_sha256':sha256(prof)})
        rec={'status':'PASS','docx_sha256':sha256(path),'profile_sha256':sha256(prof),
             'pdf':str(pdf),'pdf_sha256':sha256(pdf),'page_count':1,'read_only_copy':True,
             'text_coverage':1.0,'unknown_fonts':[],
             'pages':[{'number':1,'path':str(png),'sha256':sha256(png)}]}
        write_json(render,rec)
        write_json(visual,{'docx_sha256':sha256(path),'pdf_sha256':sha256(pdf),'reviewer':'test',
             'pages':[{'number':1,'sha256':sha256(png),'status':'PASS','notes':'Unit fixture.'}]})
        self.assertEqual(verify_release(path,prof,audit,render,visual)['status'],'PASS')
        original_pdf=pdf.read_bytes()
        canvas=Canvas(str(pdf));canvas.drawString(40,700,'Unrelated text.');canvas.save()
        bad=copy.deepcopy(rec);bad['pdf_sha256']=sha256(pdf);write_json(render,bad)
        write_json(visual,{'docx_sha256':sha256(path),'pdf_sha256':sha256(pdf),'reviewer':'test',
             'pages':[{'number':1,'sha256':sha256(png),'status':'PASS','notes':'Unit fixture.'}]})
        self.assertIn('render_text_coverage',verify_release(path,prof,audit,render,visual)['reasons'])
        pdf.write_bytes(original_pdf);write_json(render,rec)
        write_json(visual,{'docx_sha256':sha256(path),'pdf_sha256':sha256(pdf),'reviewer':'test',
             'pages':[{'number':1,'sha256':sha256(png),'status':'PASS','notes':'Unit fixture.'}]})
        doc.paragraphs[5].add_run('Changed text.');doc.save(path)
        actual=verify_release(path,prof,audit,render,visual)
        self.assertEqual(actual['status'],'BLOCK')
        self.assertIn('render_stale',actual['reasons'])

    def test_profile_declines_native_equations(self):
        from journal_profile import apply_profile
        doc,profile=self.sample()
        equation=OxmlElement('m:oMath');doc.paragraphs[5]._p.append(equation)
        with self.assertRaises(ValueError):apply_profile(doc,profile)

    def test_short_table_keeps_rows_together(self):
        from journal_profile import apply_profile
        doc,profile=self.sample();doc.add_paragraph('Table 1 Results','Caption')
        table=doc.add_table(rows=4,cols=2)
        for row in table.rows:
            for cell in row.cells:cell.text='Short comparable value'
        apply_profile(doc,profile)
        self.assertTrue(table.cell(1,0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(table.cell(3,0).paragraphs[0].paragraph_format.keep_with_next)

    def test_figure_caption_keeps_source_note(self):
        from journal_profile import apply_profile
        doc,profile=self.sample()
        cap=doc.add_paragraph('Figure 1 Operating results','Caption')
        doc.add_paragraph('Source and scope.','SourceNote')
        apply_profile(doc,profile)
        self.assertTrue(cap.paragraph_format.keep_with_next)

    def test_heading_bold_override_is_rejected(self):
        from journal_profile import audit_profile
        doc,profile=self.sample();doc.paragraphs[0].runs[0].bold=False
        self.assertIn('font_bold',[x['code'] for x in audit_profile(doc,profile)['errors']])

    def test_at_least_line_spacing_is_not_exact_spacing(self):
        from journal_profile import audit_profile
        from docx.enum.text import WD_LINE_SPACING
        doc,profile=self.sample('zh')
        doc.paragraphs[5].paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        self.assertIn('line_spacing_rule',[x['code'] for x in audit_profile(doc,profile)['errors']])

    def test_reformat_preserves_existing_column_proportions(self):
        from journal_profile import apply_profile
        doc,profile=self.sample();doc.add_paragraph('Table 1 Results','Caption')
        table=doc.add_table(rows=2,cols=2)
        apply_profile(doc,profile,{'tables':{'1':{'widths_cm':[6.0,9.92]}}})
        before=[c.width.cm for c in table.columns]
        apply_profile(doc,profile)
        for expected,actual in zip(before,[c.width.cm for c in table.columns]):
            self.assertAlmostEqual(expected,actual,places=2)


if __name__ == '__main__':
    unittest.main()
